"""Build the fixture templates the smoke test runs against.

One representative resume template (scalar header fields, a rich hyperlink,
an inline skills loop, nested experience → bullets loops, education,
a footer that reuses a header field), plus two mutations that exercise
preflight: a tag Word split across runs, and a malformed tag.

Run: .venv/bin/python scripts/make_fixtures.py
"""

from __future__ import annotations

import io
import pathlib
import zipfile

from docx import Document

FIXTURES = pathlib.Path(__file__).resolve().parent.parent / "fixtures"


def build_resume() -> bytes:
    doc = Document()
    doc.add_heading("{{ header.name }}", level=0)
    p = doc.add_paragraph()
    p.add_run("{{ header.title }}").italic = True
    doc.add_paragraph("{{ header.email }} · {{ header.phone }} · {{r header.website }}")

    doc.add_heading("Summary", level=1)
    doc.add_paragraph("{{ summary }}")

    doc.add_heading("Skills", level=1)
    doc.add_paragraph("{% for skill in skills %}{{ skill }}{% if not loop.last %} · {% endif %}{% endfor %}")

    doc.add_heading("Experience", level=1)
    doc.add_paragraph("{%p for job in experience %}")
    p = doc.add_paragraph()
    p.add_run("{{ job.title }}").bold = True
    p.add_run(" — {{ job.company }}{% if job.location %}, {{ job.location }}{% endif %}")
    doc.add_paragraph("{{ job.start }} – {{ job.end }}")
    doc.add_paragraph("{%p for bullet in job.bullets %}")
    doc.add_paragraph("{{ bullet }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p endfor %}")

    doc.add_heading("Education", level=1)
    doc.add_paragraph("{%p for school in education %}")
    p = doc.add_paragraph()
    p.add_run("{{ school.degree }}").bold = True
    p.add_run(", {{ school.institution }} ({{ school.year }})")
    doc.add_paragraph("{%p endfor %}")

    doc.sections[0].footer.paragraphs[0].text = "{{ header.name }} — references on request"

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def mutate_document_xml(docx: bytes, old: str, new: str) -> bytes:
    """Rewrite word/document.xml the way Word itself would have saved it."""
    src = zipfile.ZipFile(io.BytesIO(docx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/document.xml":
                text = data.decode("utf-8")
                assert old in text, f"{old!r} not found in document.xml"
                data = text.replace(old, new, 1).encode("utf-8")
            dst.writestr(item, data)
    return out.getvalue()


def main() -> None:
    FIXTURES.mkdir(exist_ok=True)
    resume = build_resume()
    (FIXTURES / "resume-template.docx").write_bytes(resume)

    # Word splits a tag across runs whenever spell-check or a formatting
    # change touches part of it; docxtpl's patch_xml must reassemble it.
    split = mutate_document_xml(
        resume,
        "{{ header.title }}",
        "{{ header.</w:t></w:r><w:r><w:t>title }}",
    )
    (FIXTURES / "resume-template-split-run.docx").write_bytes(split)

    # A tag the author never closed — preflight must name the paragraph.
    malformed = mutate_document_xml(resume, "{{ summary }}", "{{ summary ")
    (FIXTURES / "resume-template-malformed.docx").write_bytes(malformed)

    for path in sorted(FIXTURES.glob("*.docx")):
        print(f"{path.name}: {path.stat().st_size} bytes")


if __name__ == "__main__":
    main()

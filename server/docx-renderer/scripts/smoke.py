"""Renderer smoke: contract derivation, a full render, and the failure paths.

Runs the service in-process through FastAPI's test client so the assertions
cover the HTTP contract the Next.js side will code against, not just the
Python functions. Exit code is the verdict.

Run: .venv/bin/python scripts/smoke.py
"""

from __future__ import annotations

import io
import json
import pathlib
import sys
import zipfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

import base64  # noqa: E402

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import app  # noqa: E402

FIXTURES = pathlib.Path(__file__).resolve().parent.parent / "fixtures"
client = TestClient(app)
failures: list[str] = []


def check(condition: bool, label: str, detail: object = None) -> None:
    mark = "ok " if condition else "FAIL"
    print(f"  {mark} {label}" + (f" — {detail}" if detail is not None and not condition else ""))
    if not condition:
        failures.append(label)


def inspect(name: str) -> tuple[int, dict]:
    res = client.post("/inspect", files={"template": (name, (FIXTURES / name).read_bytes())})
    return res.status_code, res.json()


def render(name: str, values: dict, expected_checksum: str | None = None) -> tuple[int, dict]:
    data = {"values": json.dumps(values)}
    if expected_checksum:
        data["expectedChecksum"] = expected_checksum
    res = client.post("/render", files={"template": (name, (FIXTURES / name).read_bytes())}, data=data)
    return res.status_code, res.json()


VALUES = {
    "header": {
        "name": "Ada Lovelace",
        "title": "Analytical Engine Programmer & Mathematician",
        "email": "ada@example.org",
        "phone": "+44 20 7946 0000",
        "website": {"text": "ada.example.org", "url": "https://ada.example.org/?ref=cv&v=2"},
    },
    "summary": "Wrote the first published algorithm for a machine; <notes> on Babbage's Engine — 1843.",
    "skills": ["Mathematics", "Bernoulli numbers", "Algorithm design", "Technical writing"],
    "experience": [
        {
            "title": "Collaborator",
            "company": "Charles Babbage",
            "location": "London",
            "start": "1842",
            "end": "1843",
            "bullets": [
                "Translated Menabrea's memoir and tripled its length with original notes (A–G).",
                "Described how the Engine could compute Bernoulli numbers — Note G.",
            ],
        },
        {
            "title": "Independent study",
            "company": "Self-directed",
            "location": None,
            "start": "1833",
            "end": "1842",
            "bullets": ["Studied under Mary Somerville & Augustus De Morgan."],
        },
    ],
    "education": [
        {"degree": "Private tuition in mathematics", "institution": "Augustus De Morgan", "year": "1840–1842"},
    ],
}


def main() -> int:
    print("inspect resume-template.docx")
    status, contract = inspect("resume-template.docx")
    check(status == 200 and contract["ok"], "contract ok", contract)
    fields = {f["path"]: f for f in contract["fields"]}
    sections = {s["path"]: s for s in contract["sections"]}
    check({"header.name", "header.title", "header.email", "header.phone", "header.website", "summary"} <= set(fields), "scalar fields derived", sorted(fields))
    check(fields.get("header.website", {}).get("rich") is True, "{{r }} tag marked rich")
    check(fields.get("summary", {}).get("rich") is False, "plain tag not rich")
    check({"skills", "experience", "education"} == set(sections), "sections derived", sorted(sections))
    check(sections["skills"]["itemIsValue"] is True, "skills is a list of values")
    exp_fields = {f["path"]: f for f in sections["experience"]["fields"]}
    check({"title", "company", "location", "start", "end"} == set(exp_fields), "experience item fields", sorted(exp_fields))
    check(exp_fields["location"]["required"] is False, "{% if %}-guarded field is optional")
    check(exp_fields["title"]["required"] is True, "unguarded field is required")
    nested = {s["path"]: s for s in sections["experience"]["sections"]}
    check("bullets" in nested and nested["bullets"]["itemIsValue"], "nested bullets loop derived", nested)

    print("render resume-template.docx")
    status, result = render("resume-template.docx", VALUES, contract["checksum"])
    check(status == 200 and result["ok"], "render ok", result.get("error"))
    diag = result["diagnostics"]
    check(diag["unresolved"] == [], "no unresolved tags", diag["unresolved"])
    check(diag["missing"] == [], "no missing fields", diag["missing"])
    check(diag["unused"] == [], "no unused values", diag["unused"])
    text = result["text"]
    for needle in ["Ada Lovelace", "Bernoulli numbers", "<notes>", "Babbage's", "Note G", "1840–1842", "ada.example.org", "references on request"]:
        check(needle in text, f"output text contains {needle!r}")
    check("None" not in text, "null optional renders empty, not 'None'")
    check(", London" in text and "Self-directed\n" in text, "guarded location present/absent per item")
    docx = base64.b64decode(result["docxBase64"])
    doc = Document(io.BytesIO(docx))
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    check(len(bullets) == 3, "bullet paragraphs multiplied by the nested loop", bullets)
    check(not any("{%" in p.text or "{{" in p.text for p in doc.paragraphs), "no tag paragraphs survive")
    package = zipfile.ZipFile(io.BytesIO(docx))
    xml = package.read("word/document.xml").decode()
    rels = package.read("word/_rels/document.xml.rels").decode()
    check("<w:hyperlink " in xml, "hyperlink emitted for the rich field")
    check('Target="https://ada.example.org/?ref=cv&amp;v=2"' in rels, "hyperlink URL escaped in the relationship", rels[-300:])
    check("&lt;notes&gt;" in xml and "Babbage's" in xml or "Babbage&#x27;s" in xml, "angle brackets escaped in document XML")
    (FIXTURES / "out").mkdir(exist_ok=True)
    (FIXTURES / "out" / "resume-rendered.docx").write_bytes(docx)

    print("document properties")
    props = {
        "title": "David Valentine - Acme - Analyst - 20261014",
        "subject": "Analyst - Acme",
        "keywords": ["Salesforce", "HubSpot", " ", "RevOps"],
        "author": "David Valentine",
        "comments": "x" * 5000,
        "bogus": "dropped",
    }
    res = client.post(
        "/render",
        files={"template": ("t.docx", (FIXTURES / "resume-template.docx").read_bytes())},
        data={"values": json.dumps(VALUES), "properties": json.dumps(props)},
    )
    body = res.json()
    check(res.status_code == 200 and body["ok"], "render with properties ok", body.get("error"))
    written = body["diagnostics"].get("properties", {})
    check(written.get("title") == props["title"], "title set", written)
    check(written.get("subject") == "Analyst - Acme", "subject set")
    check(written.get("keywords") == "Salesforce, HubSpot, RevOps", "keyword list joined, blanks dropped", written.get("keywords"))
    check(written.get("author") == "David Valentine", "author set")
    check(len(written.get("comments", "")) == 255, "comments bounded to 255 (the OOXML limit)", len(written.get("comments", "")))
    check("bogus" not in written, "unknown property dropped")
    from render import clean_properties  # noqa: E402
    long_kw = ", ".join(f"keyword{i:03d}" for i in range(60))
    cut = clean_properties({"keywords": long_kw})["keywords"]
    check(len(cut) <= 255 and not cut.endswith(",") and cut.split(", ")[-1].startswith("keyword") and len(cut.split(", ")[-1]) == 10, "long keywords cut at a whole keyword", cut[-30:])
    core = Document(io.BytesIO(base64.b64decode(body["docxBase64"]))).core_properties
    check(core.keywords == "Salesforce, HubSpot, RevOps" and core.title == props["title"], "properties are in the file itself, not only the report")
    res = client.post(
        "/render",
        files={"template": ("t.docx", (FIXTURES / "resume-template.docx").read_bytes())},
        data={"values": json.dumps(VALUES), "properties": "{nope"},
    )
    check(res.status_code == 400 and res.json()["detail"]["code"] == "properties_not_json", "malformed properties refused", res.json())

    print("render with a missing field and a typo'd key")
    partial = {**VALUES, "sumary": VALUES["summary"]}
    del partial["summary"]
    status, result = render("resume-template.docx", partial)
    check(status == 200 and result["ok"], "render still succeeds")
    check(result["diagnostics"]["missing"] == ["summary"], "missing field reported", result["diagnostics"]["missing"])
    check(result["diagnostics"]["unused"] == ["sumary"], "unused key reported", result["diagnostics"]["unused"])

    print("checksum mismatch")
    status, result = render("resume-template.docx", VALUES, "0" * 64)
    check(status == 422 and result["error"]["code"] == "template_checksum_mismatch", "stale contract refused", result)

    print("inspect resume-template-split-run.docx")
    status, contract = inspect("resume-template-split-run.docx")
    check(status == 200 and contract["ok"], "split-run tag reassembled", contract.get("errors"))
    check("header.title" in {f["path"] for f in contract["fields"]}, "split field still derived")
    status, result = render("resume-template-split-run.docx", VALUES)
    check(result["ok"] and "Analytical Engine Programmer" in result["text"], "split-run tag renders")

    print("inspect resume-template-malformed.docx")
    status, contract = inspect("resume-template-malformed.docx")
    check(status == 422 and not contract["ok"], "malformed tag rejected")
    err = (contract["errors"] or [{}])[0]
    check(err.get("code") == "template_syntax" and err.get("paragraph"), "error names a paragraph", err)
    check("summary" in (err.get("near") or ""), "error quotes the paragraph text", err)
    status, result = render("resume-template-malformed.docx", VALUES)
    check(status == 422 and result["error"]["code"] == "template_invalid", "render refuses a template that fails preflight")

    print("request guards")
    res = client.post("/render", files={"template": ("x.docx", b"not a zip")}, data={"values": "{}"})
    check(res.status_code == 422 and res.json()["error"]["code"] == "template_invalid", "non-docx bytes rejected", res.json())
    res = client.post("/render", files={"template": ("x.docx", (FIXTURES / "resume-template.docx").read_bytes())}, data={"values": "[1]"})
    check(res.status_code == 400, "values must be an object")
    res = client.post("/render", files={"template": ("x.docx", (FIXTURES / "resume-template.docx").read_bytes())}, data={"values": "{nope"})
    check(res.status_code == 400, "values must be JSON")

    print()
    if failures:
        print(f"{len(failures)} failure(s):", *failures, sep="\n  ")
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
